import csv
import io
import uuid
from pathlib import Path

from fastapi import HTTPException, UploadFile, status
from pdfminer.high_level import extract_pages
from pdfminer.layout import LTTextContainer
from docx import Document as DocxDocument
from pptx import Presentation
from sqlalchemy.orm import Session

from app.models.chunk import Chunk
from app.models.document import Document
from app.models.user import User
from app.schemas.document import UploadResponse
from app.services.retrieval_service import retrieval_service
from app.services.storage_service import storage_service
from app.services.workspace_service import get_workspace_for_user


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".txt", ".md", ".markdown", ".csv"}


def _chunk_text(text: str, size: int = 1200, overlap: int = 200) -> list[str]:
    if not text.strip():
        return []
    chunks: list[str] = []
    start = 0
    while start < len(text):
        end = min(start + size, len(text))
        chunks.append(text[start:end].strip())
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return [chunk for chunk in chunks if chunk]


def _extract_text(path: str) -> tuple[str, list[tuple[int | None, str]]]:
    suffix = Path(path).suffix.lower()
    file_bytes = Path(path).read_bytes()

    if suffix in {".txt", ".md", ".markdown"}:
        text = file_bytes.decode("utf-8", errors="ignore")
        return text, [(None, text)]
    if suffix == ".csv":
        decoded = file_bytes.decode("utf-8", errors="ignore")
        reader = csv.reader(io.StringIO(decoded))
        rows = [" | ".join(row) for row in reader]
        text = "\n".join(rows)
        return text, [(None, text)]
    if suffix == ".pdf":
        pages: list[tuple[int | None, str]] = []
        for page_number, page_layout in enumerate(extract_pages(path), start=1):
            fragments: list[str] = []
            for element in page_layout:
                if isinstance(element, LTTextContainer):
                    fragments.append(element.get_text())
            page_text = "\n".join(fragment.strip() for fragment in fragments if fragment.strip())
            if page_text:
                pages.append((page_number, page_text))
        full_text = "\n".join(text for _, text in pages)
        return full_text, pages
    if suffix == ".docx":
        document = DocxDocument(path)
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        text = "\n".join(paragraphs)
        return text, [(1, text)]
    if suffix == ".pptx":
        presentation = Presentation(path)
        pages: list[tuple[int | None, str]] = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            slide_text: list[str] = []
            for shape in slide.shapes:
                text = getattr(shape, "text", "").strip()
                if text:
                    slide_text.append(text)
            if slide_text:
                pages.append((slide_number, "\n".join(slide_text)))
        full_text = "\n".join(text for _, text in pages)
        return full_text, pages

    text = file_bytes.decode("utf-8", errors="ignore")
    return text, [(1, text)]


async def ingest_document(
    workspace_id: str,
    file: UploadFile,
    user: User,
    db: Session,
) -> UploadResponse:
    get_workspace_for_user(workspace_id, user, db)

    suffix = Path(file.filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Unsupported file type: {suffix}",
        )

    storage_path = await storage_service.save_document(workspace_id, file)
    document = Document(
        workspace_id=workspace_id,
        filename=file.filename,
        source_type="upload",
        mime_type=file.content_type,
        storage_path=storage_path,
        status="processing",
    )
    db.add(document)
    db.commit()
    db.refresh(document)

    _, pages = _extract_text(storage_path)
    if not any(text.strip() for _, text in pages):
        document.status = "failed"
        db.commit()
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="No readable text could be extracted from this file.",
        )
    total_chunks = 0
    for page_number, text in pages:
        for index, content in enumerate(_chunk_text(text)):
            chunk = Chunk(
                workspace_id=workspace_id,
                document_id=document.id,
                chunk_id=f"{document.id}:{page_number or 0}:{index}:{uuid.uuid4().hex[:8]}",
                page_number=page_number,
                content=content,
                source_type="document",
            )
            db.add(chunk)
            db.flush()
            retrieval_service.index_chunk(chunk, document.filename)
            total_chunks += 1

    document.status = "ready"
    db.commit()
    db.refresh(document)
    return UploadResponse(document=document, ingested_chunks=total_chunks)


def list_documents(workspace_id: str, user: User, db: Session) -> list[Document]:
    get_workspace_for_user(workspace_id, user, db)
    return (
        db.query(Document)
        .filter(Document.workspace_id == workspace_id)
        .order_by(Document.created_at.desc())
        .all()
    )
